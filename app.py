import sqlite3
import os
import io
import uuid
import smtplib
import threading
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from flask import Flask, render_template, request, redirect
from flask import session, send_from_directory, send_file
from flask_socketio import SocketIO, emit
from markupsafe import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.security import generate_password_hash
from werkzeug.security import check_password_hash

app = Flask(__name__)
socketio = SocketIO(app)
app.secret_key = os.environ.get(
    "SECRET_KEY", "change-this-to-any-long-random-text"
)

UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp"}

DEPARTMENTS = [
    "IT", "Bio Medical", "HR", "Accounts", "Swasta Ward",
    "ICU", "NICU", "Management", "Admin", "Lab",
    "Radiology", "Security", "Canteen", "OPD", "Billing",
    "Insurance", "Facility",
]

IST = timezone(timedelta(hours=5, minutes=30))

DEPT_BAR_CSS = (
    "<style>"
    ".dept-bar{position:fixed;top:0;left:0;right:0;z-index:1000;"
    "background:#8b1e1e;color:#ffffff;padding:10px 16px;"
    "font-size:15px;font-weight:600;text-align:center;"
    "border-bottom:3px solid #d4a017;"
    "box-shadow:0 2px 8px rgba(0,0,0,0.25);}"
    ".dept-bar b{color:#f5c542;}"
    "body{padding-top:70px !important;}"
    "</style>"
)


def now_ist():
    return datetime.now(IST).strftime("%Y-%m-%d %H:%M:%S")


def split_dt(value):
    if not value:
        return "", ""
    parts = value.split(" ")
    if len(parts) == 2:
        return parts[0], parts[1]
    return value, ""


def allowed_file(filename):
    if "." not in filename:
        return False
    ext = filename.rsplit(".", 1)[1].lower()
    return ext in ALLOWED_EXTENSIONS


def get_db():
    conn = sqlite3.connect("data.db")
    conn.row_factory = sqlite3.Row
    return conn


def add_column(conn, sql):
    try:
        conn.execute(sql)
    except sqlite3.OperationalError:
        pass


def init_db():
    conn = get_db()
    conn.execute(
        "CREATE TABLE IF NOT EXISTS names ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT, "
        "name TEXT NOT NULL)"
    )
    conn.execute(
        "CREATE TABLE IF NOT EXISTS tickets ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT, "
        "title TEXT NOT NULL, "
        "description TEXT NOT NULL, "
        "priority TEXT NOT NULL, "
        "status TEXT NOT NULL DEFAULT 'Open', "
        "created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)"
    )
    conn.execute(
        "CREATE TABLE IF NOT EXISTS users ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT, "
        "username TEXT NOT NULL UNIQUE, "
        "password_hash TEXT NOT NULL, "
        "role TEXT NOT NULL DEFAULT 'Staff')"
    )
    conn.execute(
        "CREATE TABLE IF NOT EXISTS ticket_depts ("
        "ticket_id INTEGER NOT NULL, "
        "dept TEXT NOT NULL, "
        "PRIMARY KEY (ticket_id, dept))"
    )
    add_column(conn, "ALTER TABLE tickets ADD COLUMN image TEXT")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN assigned_to INTEGER")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN raised_by TEXT")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN closed_by TEXT")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN closed_at TEXT")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN raised_by_name TEXT")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN handled_by_name TEXT")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN handled_by_dept TEXT")
    add_column(conn, "ALTER TABLE tickets ADD COLUMN closed_by_name TEXT")
    add_column(conn, "ALTER TABLE users ADD COLUMN email TEXT")
    conn.commit()
    conn.close()


def seed_departments():
    default_pw = os.environ.get("DEFAULT_PASSWORD", "Welcome@123")
    conn = get_db()
    for dept in DEPARTMENTS:
        row = conn.execute(
            "SELECT id FROM users WHERE username = ?", (dept,)
        ).fetchone()
        if not row:
            role = "Admin" if dept == "Admin" else "Dept"
            conn.execute(
                "INSERT INTO users (username, password_hash, role) "
                "VALUES (?, ?, ?)",
                (dept, generate_password_hash(default_pw), role),
            )
    conn.commit()
    conn.close()


init_db()
seed_departments()


@app.after_request
def add_department_bar(response):
    if response.mimetype != "text/html":
        return response
    if response.direct_passthrough:
        return response
    name = session.get("username")
    if not name:
        return response
    html = response.get_data(as_text=True)
    if "<body>" not in html:
        return response
    bar = (
        DEPT_BAR_CSS
        + '<div class="dept-bar">Department: <b>'
        + str(escape(name))
        + "</b></div>"
    )
    html = html.replace("<body>", "<body>" + bar, 1)
    response.set_data(html)
    return response


def _send_email(to_list, subject, body):
    sender = os.environ.get("GMAIL_USER")
    password = os.environ.get("GMAIL_APP_PASSWORD")
    if not sender or not password or not to_list:
        print("Email skipped: variables not set or no recipients.", flush=True)
        return
    try:
        msg = EmailMessage()
        msg["From"] = sender
        msg["To"] = ", ".join(to_list)
        msg["Subject"] = subject
        msg.set_content(body)
        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=20) as server:
            server.login(sender, password)
            server.send_message(msg)
        print("Email sent to:", to_list, flush=True)
    except Exception as e:
        print("Email failed:", e, flush=True)


def notify_admins(subject, body):
    conn = get_db()
    rows = conn.execute(
        "SELECT email FROM users "
        "WHERE role = 'Admin' AND email IS NOT NULL AND email != ''"
    ).fetchall()
    conn.close()
    to_list = [r["email"] for r in rows]
    t = threading.Thread(target=_send_email, args=(to_list, subject, body))
    t.daemon = True
    t.start()


def clean_depts(values):
    good = [d for d in values if d in DEPARTMENTS]
    return list(dict.fromkeys(good))


def can_access(conn, ticket_id):
    if session.get("role") == "Admin":
        return True
    me = session.get("username")
    row = conn.execute(
        "SELECT 1 FROM tickets WHERE id = ? AND (raised_by = ? OR id IN "
        "(SELECT ticket_id FROM ticket_depts WHERE dept = ?))",
        (ticket_id, me, me),
    ).fetchone()
    return row is not None


def is_assigned(conn, ticket_id):
    row = conn.execute(
        "SELECT 1 FROM ticket_depts WHERE ticket_id = ? AND dept = ?",
        (ticket_id, session.get("username")),
    ).fetchone()
    return row is not None


def load_tickets():
    conn = get_db()
    if session.get("role") == "Admin":
        rows = conn.execute(
            "SELECT * FROM tickets ORDER BY id DESC"
        ).fetchall()
    else:
        me = session.get("username")
        rows = conn.execute(
            "SELECT * FROM tickets WHERE raised_by = ? OR id IN "
            "(SELECT ticket_id FROM ticket_depts WHERE dept = ?) "
            "ORDER BY id DESC",
            (me, me),
        ).fetchall()
    result = []
    for r in rows:
        t = dict(r)
        d = conn.execute(
            "SELECT dept FROM ticket_depts WHERE ticket_id = ? "
            "ORDER BY dept",
            (t["id"],),
        ).fetchall()
        t["depts"] = [x["dept"] for x in d]
        result.append(t)
    conn.close()
    return result


@app.route("/")
def home():
    return render_template("home.html")


@app.route("/about")
def about():
    return "This is my about page."


@app.route("/form", methods=["GET", "POST"])
def form():
    name = None
    if request.method == "POST":
        name = request.form["username"]
        conn = get_db()
        conn.execute("INSERT INTO names (name) VALUES (?)", (name,))
        conn.commit()
        conn.close()
    return render_template("form.html", name=name)


@app.route("/names")
def names():
    conn = get_db()
    rows = conn.execute("SELECT * FROM names ORDER BY id DESC").fetchall()
    conn.close()
    return render_template("names.html", rows=rows)


@app.route("/chat")
def chat():
    return render_template("chat.html")


@app.route("/tickets/new", methods=["GET", "POST"])
def new_ticket():
    if "user_id" not in session:
        return redirect("/login")
    if request.method == "POST":
        title = request.form["title"]
        description = request.form["description"]
        priority = request.form["priority"]
        person = request.form.get("your_name", "").strip()
        if not person:
            return "Please enter your name.", 400
        depts = clean_depts(request.form.getlist("depts"))

        image_name = None
        file = request.files.get("camera_image")
        if not file or file.filename == "":
            file = request.files.get("gallery_image")
        if file and file.filename != "" and allowed_file(file.filename):
            ext = file.filename.rsplit(".", 1)[1].lower()
            image_name = uuid.uuid4().hex + "." + ext
            file.save(os.path.join(UPLOAD_FOLDER, image_name))

        who = session.get("username", "unknown")
        conn = get_db()
        cur = conn.execute(
            "INSERT INTO tickets "
            "(title, description, priority, image, raised_by, "
            "raised_by_name, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (title, description, priority, image_name, who,
             person, now_ist()),
        )
        ticket_id = cur.lastrowid
        for d in depts:
            conn.execute(
                "INSERT OR IGNORE INTO ticket_depts (ticket_id, dept) "
                "VALUES (?, ?)",
                (ticket_id, d),
            )
        conn.commit()
        conn.close()
        socketio.emit("new_ticket_alert", {"title": title})
        socketio.emit("tickets_changed")

        body = "A new ticket was raised by " + person
        body += " (" + who + ").\n\n"
        body += "Title: " + title + "\n"
        body += "Priority: " + priority + "\n\n"
        body += "Description:\n" + description + "\n"
        notify_admins("New ticket: " + title, body)

        return render_template(
            "ticket_new.html", saved=True, departments=DEPARTMENTS
        )
    return render_template(
        "ticket_new.html", saved=False, departments=DEPARTMENTS
    )


@app.route("/tickets")
def tickets():
    if "user_id" not in session:
        return redirect("/login")
    rows = load_tickets()
    return render_template(
        "tickets.html", rows=rows, departments=DEPARTMENTS
    )


@app.route("/tickets/export")
def export_tickets():
    if "user_id" not in session:
        return redirect("/login")
    rows = load_tickets()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    headers = [
        "ID", "Title", "Description", "Priority", "Status",
        "Raised By (Dept)", "Raised By (Person)", "Assigned To",
        "Raised Date", "Raised Time",
        "Last Handled By", "Closed By (Dept)", "Closed By (Person)",
        "Closed Date", "Closed Time",
    ]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="8B1E1E")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for t in rows:
        r_date, r_time = split_dt(t.get("created_at"))
        c_date, c_time = split_dt(t.get("closed_at"))
        handled = ""
        if t.get("handled_by_name"):
            handled = t["handled_by_name"]
            if t.get("handled_by_dept"):
                handled += " (" + t["handled_by_dept"] + ")"
        ws.append([
            t["id"], t["title"], t["description"], t["priority"],
            t["status"], t.get("raised_by") or "",
            t.get("raised_by_name") or "",
            ", ".join(t["depts"]), r_date, r_time,
            handled, t.get("closed_by") or "",
            t.get("closed_by_name") or "", c_date, c_time,
        ])
    widths = [6, 28, 40, 10, 12, 16, 18, 30, 14, 12, 26, 16, 18, 14, 12]
    for i, w in enumerate(widths):
        col = chr(ord("A") + i)
        ws.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    stamp = datetime.now(IST).strftime("%Y%m%d_%H%M")
    return send_file(
        out,
        as_attachment=True,
        download_name="tickets_" + stamp + ".xlsx",
        mimetype="application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet",
    )


@app.route("/tickets/<int:ticket_id>/status", methods=["POST"])
def update_status(ticket_id):
    if "user_id" not in session:
        return redirect("/login")
    new_status = request.form["status"]
    person = request.form.get("person_name", "").strip()
    if not person:
        return "Please enter your name.", 400
    conn = get_db()
    allowed = session.get("role") == "Admin" or is_assigned(conn, ticket_id)
    if not allowed:
        conn.close()
        return "Only Admin or assigned department can change status.", 403
    dept = session["username"]
    if new_status in ("Open", "In Progress", "Closed"):
        if new_status == "Closed":
            conn.execute(
                "UPDATE tickets SET status = ?, handled_by_name = ?, "
                "handled_by_dept = ?, closed_by = ?, closed_by_name = ?, "
                "closed_at = ? WHERE id = ?",
                (new_status, person, dept, dept, person,
                 now_ist(), ticket_id),
            )
        else:
            conn.execute(
                "UPDATE tickets SET status = ?, handled_by_name = ?, "
                "handled_by_dept = ?, closed_by = NULL, "
                "closed_by_name = NULL, closed_at = NULL WHERE id = ?",
                (new_status, person, dept, ticket_id),
            )
        conn.commit()
    conn.close()
    socketio.emit("tickets_changed")
    return redirect("/tickets")


@app.route("/tickets/<int:ticket_id>/assign", methods=["POST"])
def assign_ticket(ticket_id):
    if "user_id" not in session:
        return redirect("/login")
    conn = get_db()
    if not can_access(conn, ticket_id):
        conn.close()
        return "You cannot assign this ticket.", 403
    depts = clean_depts(request.form.getlist("depts"))
    conn.execute(
        "DELETE FROM ticket_depts WHERE ticket_id = ?", (ticket_id,)
    )
    for d in depts:
        conn.execute(
            "INSERT OR IGNORE INTO ticket_depts (ticket_id, dept) "
            "VALUES (?, ?)",
            (ticket_id, d),
        )
    conn.commit()
    conn.close()
    socketio.emit("tickets_changed")
    return redirect("/tickets")


@app.route("/tickets/<int:ticket_id>/delete", methods=["POST"])
def delete_ticket(ticket_id):
    if session.get("role") != "Admin":
        return "Only Admin can delete tickets.", 403
    conn = get_db()
    row = conn.execute(
        "SELECT image FROM tickets WHERE id = ?", (ticket_id,)
    ).fetchone()
    if row and row["image"]:
        try:
            os.remove(os.path.join(UPLOAD_FOLDER, row["image"]))
        except OSError:
            pass
    conn.execute(
        "DELETE FROM ticket_depts WHERE ticket_id = ?", (ticket_id,)
    )
    conn.execute("DELETE FROM tickets WHERE id = ?", (ticket_id,))
    conn.commit()
    conn.close()
    socketio.emit("tickets_changed")
    return redirect("/tickets")


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        conn = get_db()
        user = conn.execute(
            "SELECT * FROM users WHERE lower(username) = lower(?)",
            (username,),
        ).fetchone()
        conn.close()
        if user and check_password_hash(user["password_hash"], password):
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user["role"]
            return redirect("/tickets")
        error = "Wrong department name or password."
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/login")


@app.route("/uploads/<filename>")
def uploaded_file(filename):
    if "user_id" not in session:
        return redirect("/login")
    return send_from_directory(UPLOAD_FOLDER, filename)


@socketio.on("send_message")
def handle_message(data):
    emit("new_message", data, broadcast=True)


from passwords import bp
app.register_blueprint(bp)


if __name__ == "__main__":
    socketio.run(app, debug=True, allow_unsafe_werkzeug=True)